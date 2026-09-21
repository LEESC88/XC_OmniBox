import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import pygments
from pygments.lexers import get_lexer_by_name
from pygments.token import Token

def create_document():
    doc = docx.Document()

    # Base Font Helper
    def set_run_font(run, font_name="Times New Roman", size_pt=12, bold=False, italic=False, color_rgb=None, east_asia="Microsoft YaHei"):
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        run.bold = bold
        run.italic = italic
        if color_rgb:
            run.font.color.rgb = color_rgb
        rPr = run._r.get_or_add_rPr()
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}" w:eastAsia="{east_asia}"/>')
        rPr.append(rFonts)

    # Page Margins & Header/Footer Setup
    for section in doc.sections:
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        
        # Remove Chinese document grid pitch that forces 18pt intervals!
        sectPr = section._sectPr
        docGrid = sectPr.find(qn('w:docGrid'))
        if docGrid is not None:
            sectPr.remove(docGrid)
        newDocGrid = parse_xml(f'<w:docGrid {nsdecls("w")} w:type="default"/>')
        sectPr.append(newDocGrid)
        
        # Header setup
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.paragraph_format.space_before = Pt(0)
        hp.paragraph_format.space_after = Pt(2)
        hrun = hp.add_run("XC_OmniBox · 桌面客户端构建与自动更新指南")
        set_run_font(hrun, font_name="Times New Roman", size_pt=8.5, color_rgb=RGBColor(140, 145, 155), east_asia="Microsoft YaHei")
        
        # Footer setup
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.space_before = Pt(3)
        fp.paragraph_format.space_after = Pt(0)
        frun = fp.add_run("— 工业级 Windows 桌面端交付方案 —")
        set_run_font(frun, font_name="Times New Roman", size_pt=8.5, color_rgb=RGBColor(160, 165, 175), east_asia="Microsoft YaHei")

    def add_title(text, subtitle=""):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.keep_with_next = True
        
        run = p.add_run(text)
        set_run_font(run, font_name="Times New Roman", size_pt=18, bold=True, color_rgb=RGBColor(15, 44, 89), east_asia="Microsoft YaHei")
        
        if subtitle:
            p_sub = doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_sub.paragraph_format.space_before = Pt(1)
            p_sub.paragraph_format.space_after = Pt(4)
            p_sub.paragraph_format.line_spacing = 1.15
            p_sub.paragraph_format.keep_with_next = True
            run_sub = p_sub.add_run(subtitle)
            set_run_font(run_sub, font_name="Times New Roman", size_pt=10.5, bold=False, italic=False, color_rgb=RGBColor(71, 85, 105), east_asia="Microsoft YaHei")

    def add_meta_bar():
        tbl = doc.add_table(rows=1, cols=3)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        widths = [Inches(2.3), Inches(2.4), Inches(2.3)]
        
        items = [
            ("📁 项目名称", "XC_OmniBox (万象箱)"),
            ("⚙️ 核心架构", "Electron + Next.js + FastAPI"),
            ("🚀 分发交付", "Windows NSIS 静默更新安装包")
        ]
        
        tblPr = tbl._tbl.tblPr
        tblBorders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
                <w:insideH w:val="none"/>
                <w:insideV w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            </w:tblBorders>
        ''')
        tblPr.append(tblBorders)
        
        tblCellMar = parse_xml(f'''
            <w:tblCellMar {nsdecls("w")}>
                <w:top w:w="50" w:type="dxa"/>
                <w:left w:w="120" w:type="dxa"/>
                <w:bottom w:w="50" w:type="dxa"/>
                <w:right w:w="120" w:type="dxa"/>
            </w:tblCellMar>
        ''')
        tblPr.append(tblCellMar)
        
        trPr = tbl.rows[0]._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        for idx, cell in enumerate(tbl.rows[0].cells):
            cell.width = widths[idx]
            tcPr = cell._tc.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
            tcPr.append(shd)
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            label, val = items[idx]
            r1 = p.add_run(f"{label}: ")
            set_run_font(r1, font_name="Times New Roman", size_pt=9, bold=True, color_rgb=RGBColor(51, 65, 85), east_asia="Microsoft YaHei")
            r2 = p.add_run(val)
            set_run_font(r2, font_name="Times New Roman", size_pt=9, bold=False, color_rgb=RGBColor(15, 23, 42), east_asia="Microsoft YaHei")

    def add_h1(text, space_before=8, space_after=2):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.keep_with_next = True
        
        run = p.add_run(text)
        set_run_font(run, font_name="Times New Roman", size_pt=13.5, bold=True, color_rgb=RGBColor(15, 44, 89), east_asia="Microsoft YaHei")
        return p

    def add_h2(text, space_before=5, space_after=1.5):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.keep_with_next = True
        
        run = p.add_run(text)
        set_run_font(run, font_name="Times New Roman", size_pt=12, bold=True, color_rgb=RGBColor(30, 41, 59), east_asia="Microsoft YaHei")
        return p

    def add_body(text_runs, space_after=2, space_before=0, line_spacing=1.5, keep_with_next=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        p.paragraph_format.keep_with_next = keep_with_next
        
        if isinstance(text_runs, str):
            text_runs = [(text_runs, False, False, RGBColor(38, 38, 38))]
            
        for item in text_runs:
            txt = item[0]
            bold = item[1] if len(item) > 1 else False
            italic = item[2] if len(item) > 2 else False
            color = item[3] if len(item) > 3 and item[3] else RGBColor(38, 38, 38)
            run = p.add_run(txt)
            set_run_font(run, font_name="Times New Roman", size_pt=12, bold=bold, italic=italic, color_rgb=color, east_asia="Microsoft YaHei")
        return p

    def add_bullet(runs, level=0, space_after=2):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22 * (level + 1))
        p.paragraph_format.first_line_indent = Inches(-0.16)
        p.paragraph_format.space_before = Pt(0.5)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.5
        
        bullet_sym = "• " if level == 0 else "▫ "
        r_sym = p.add_run(bullet_sym)
        set_run_font(r_sym, font_name="Times New Roman", size_pt=12, bold=True, color_rgb=RGBColor(15, 44, 89), east_asia="Microsoft YaHei")

        if isinstance(runs, str):
            runs = [(runs, False, False, RGBColor(38, 38, 38))]
            
        for item in runs:
            txt = item[0]
            bold = item[1] if len(item) > 1 else False
            italic = item[2] if len(item) > 2 else False
            color = item[3] if len(item) > 3 and item[3] else RGBColor(38, 38, 38)
            run = p.add_run(txt)
            set_run_font(run, font_name="Times New Roman", size_pt=12, bold=bold, italic=italic, color_rgb=color, east_asia="Microsoft YaHei")
        return p

    def add_callout(title, items, box_type="info"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.columns[0].width = Inches(7.0)
        
        cell = tbl.rows[0].cells[0]
        tcPr = cell._tc.get_or_add_tcPr()
        
        bg_color = "F0FDF4" if box_type == "success" else ("FEF9C3" if box_type == "warning" else "F8FAFC")
        border_color = "16A34A" if box_type == "success" else ("CA8A04" if box_type == "warning" else "2563EB")
        
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
        tcPr.append(shd)
        
        tblPr = tbl._tbl.tblPr
        tblBorders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:bottom w:val="none"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
                <w:right w:val="none"/>
                <w:insideH w:val="none"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr.append(tblBorders)
        
        tblCellMar = parse_xml(f'''
            <w:tblCellMar {nsdecls("w")}>
                <w:top w:w="50" w:type="dxa"/>
                <w:left w:w="120" w:type="dxa"/>
                <w:bottom w:w="50" w:type="dxa"/>
                <w:right w:w="120" w:type="dxa"/>
            </w:tblCellMar>
        ''')
        tblPr.append(tblCellMar)
        
        trPr = tbl.rows[0]._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.keep_with_next = True
        
        r_title = p.add_run(title)
        title_color = RGBColor(22, 101, 52) if box_type == "success" else (RGBColor(133, 77, 14) if box_type == "warning" else RGBColor(30, 64, 175))
        set_run_font(r_title, font_name="Times New Roman", size_pt=10.5, bold=True, color_rgb=title_color, east_asia="Microsoft YaHei")
        
        for idx, item in enumerate(items):
            p_item = cell.add_paragraph()
            p_item.paragraph_format.space_before = Pt(1)
            p_item.paragraph_format.space_after = Pt(1)
            p_item.paragraph_format.line_spacing = 1.3
            if idx < len(items) - 1:
                p_item.paragraph_format.keep_with_next = True
            
            if isinstance(item, str):
                item = [(item, False, False, RGBColor(51, 65, 85))]
            for r_info in item:
                txt = r_info[0]
                bld = r_info[1] if len(r_info) > 1 else False
                itl = r_info[2] if len(r_info) > 2 else False
                clr = r_info[3] if len(r_info) > 3 and r_info[3] else RGBColor(51, 65, 85)
                r = p_item.add_run(txt)
                set_run_font(r, font_name="Times New Roman", size_pt=10, bold=bld, italic=itl, color_rgb=clr, east_asia="Microsoft YaHei")

    def map_token_to_style(ttype, val, lang):
        C_KEYWORD = RGBColor(86, 156, 214)     # #569CD6
        C_CONTROL = RGBColor(197, 134, 192)    # #C586C0
        C_STRING = RGBColor(206, 145, 120)     # #CE9178
        C_NUMBER = RGBColor(181, 206, 168)     # #B5CEA8
        C_COMMENT = RGBColor(106, 153, 85)     # #6A9955
        C_FUNC = RGBColor(220, 220, 170)       # #DCDCAA
        C_VAR = RGBColor(156, 220, 254)        # #9CDCFE
        C_TYPE = RGBColor(78, 201, 176)        # #4EC9B0
        C_DEFAULT = RGBColor(212, 212, 212)    # #D4D4D4
        C_CONSTANT = RGBColor(79, 193, 255)    # #4FC1FF

        if ttype in Token.Comment:
            return C_COMMENT, False, True
        if ttype in Token.Literal.String:
            return C_STRING, False, False
        if ttype in Token.Literal.Number:
            return C_NUMBER, False, False

        val_strip = val.strip()

        if lang == 'bash':
            if val_strip.startswith('-'):
                return C_VAR, False, False
            if val_strip in ['npm', 'pip', 'pyinstaller', 'cd', 'concurrently', 'wait-on', 'electron', 'electron-builder']:
                return C_TYPE, True, False
            if val_strip in ['install', 'run', 'dist', 'dev:electron', 'build:frontend', 'build:backend']:
                return C_FUNC, False, False
            return C_DEFAULT, False, False

        if lang == 'json':
            if ttype in Token.Name.Tag:
                return C_VAR, False, False
            if val_strip in ['true', 'false', 'null']:
                return C_KEYWORD, True, False
            return C_DEFAULT, False, False

        if lang in ['javascript', 'typescript']:
            if val_strip in ['if', 'else', 'return', 'try', 'catch', 'finally', 'throw', 'for', 'while', 'switch', 'case']:
                return C_CONTROL, True, False
            if val_strip in ['const', 'let', 'var', 'function', 'require', 'export', 'default', 'import', 'from', 'new', 'async', 'await', 'class', 'type', 'typeof', 'as', 'any']:
                return C_KEYWORD, True, False
            if val_strip in ['app', 'BrowserWindow', 'ipcMain', 'ipcRenderer', 'contextBridge', 'autoUpdater', 'process', '__dirname', 'window', 'module', 'exports', 'String', 'Boolean']:
                return C_TYPE, False, False
            if val_strip in ['startBackend', 'createWindow', 'exposeInMainWorld', 'join', 'spawn', 'loadURL', 'loadFile', 'openDevTools', 'whenReady', 'on', 'setupAutoUpdater', 'checkForUpdates', 'downloadUpdate', 'quitAndInstall', 'handle', 'send', 'invoke', 'treeKill']:
                return C_FUNC, False, False
            if ttype in Token.Name.Function:
                return C_FUNC, False, False
            if val_strip.isupper() and len(val_strip) > 1:
                return C_CONSTANT, False, False
            if ttype in Token.Name:
                return C_VAR, False, False
            return C_DEFAULT, False, False

        return C_DEFAULT, False, False

    def add_vscode_block(lang, title_bar, code_str, line_height_pt=9.6, code_font_size=8.3):
        tbl = doc.add_table(rows=2, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.columns[0].width = Inches(7.0)
        
        # Outer table borders
        tblPr = tbl._tbl.tblPr
        tblBorders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="6" w:space="0" w:color="3C3C3C"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="3C3C3C"/>
                <w:left w:val="single" w:sz="6" w:space="0" w:color="3C3C3C"/>
                <w:right w:val="single" w:sz="6" w:space="0" w:color="3C3C3C"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="2D2D2D"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr.append(tblBorders)
        
        tblCellMar = parse_xml(f'''
            <w:tblCellMar {nsdecls("w")}>
                <w:top w:w="30" w:type="dxa"/>
                <w:left w:w="120" w:type="dxa"/>
                <w:bottom w:w="30" w:type="dxa"/>
                <w:right w:w="120" w:type="dxa"/>
            </w:tblCellMar>
        ''')
        tblPr.append(tblCellMar)
        
        # Row 0: Title Bar
        r0 = tbl.rows[0].cells[0]
        r0_tcPr = r0._tc.get_or_add_tcPr()
        shd0 = parse_xml(f'<w:shd {nsdecls("w")} w:fill="252526"/>')
        r0_tcPr.append(shd0)
        
        tr0_Pr = tbl.rows[0]._tr.get_or_add_trPr()
        tr0_Pr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        p0 = r0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(1.5)
        p0.paragraph_format.space_after = Pt(1.5)
        p0.paragraph_format.line_spacing = 1.15
        p0.paragraph_format.keep_with_next = True
        
        # Window buttons
        r_red = p0.add_run("● ")
        set_run_font(r_red, font_name="Segoe UI", size_pt=7, bold=True, color_rgb=RGBColor(255, 95, 86))
        r_yel = p0.add_run("● ")
        set_run_font(r_yel, font_name="Segoe UI", size_pt=7, bold=True, color_rgb=RGBColor(255, 189, 46))
        r_grn = p0.add_run("●   ")
        set_run_font(r_grn, font_name="Segoe UI", size_pt=7, bold=True, color_rgb=RGBColor(39, 201, 63))
        
        r_title = p0.add_run(title_bar)
        set_run_font(r_title, font_name="Consolas", size_pt=8, bold=True, color_rgb=RGBColor(180, 190, 205), east_asia="Microsoft YaHei")
        
        # Row 1: Code Body
        r1 = tbl.rows[1].cells[0]
        r1_tcPr = r1._tc.get_or_add_tcPr()
        shd1 = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E1E1E"/>')
        r1_tcPr.append(shd1)
        
        lines_count = len(code_str.strip().split('\n'))
        tr1_Pr = tbl.rows[1]._tr.get_or_add_trPr()
        tr1_Pr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        lexer = get_lexer_by_name(lang)
        tokens = list(pygments.lex(code_str.strip(), lexer))
        
        p_cur = r1.paragraphs[0]
        p_cur.paragraph_format.space_before = Pt(0)
        p_cur.paragraph_format.space_after = Pt(0)
        p_cur.paragraph_format.line_spacing = Pt(line_height_pt)
        pPr = p_cur._p.get_or_add_pPr()
        pPr.append(parse_xml(f'<w:snapToGrid {nsdecls("w")} w:val="0"/>'))
        
        all_paras = [p_cur]

        for ttype, val in tokens:
            if '\n' in val:
                parts = val.split('\n')
                for i, part in enumerate(parts):
                    if i > 0:
                        p_cur = r1.add_paragraph()
                        p_cur.paragraph_format.space_before = Pt(0)
                        p_cur.paragraph_format.space_after = Pt(0)
                        p_cur.paragraph_format.line_spacing = Pt(line_height_pt)
                        pPr = p_cur._p.get_or_add_pPr()
                        pPr.append(parse_xml(f'<w:snapToGrid {nsdecls("w")} w:val="0"/>'))
                        all_paras.append(p_cur)
                    if part:
                        color, bold, italic = map_token_to_style(ttype, part, lang)
                        r = p_cur.add_run(part)
                        set_run_font(r, font_name="Consolas", size_pt=code_font_size, bold=bold, italic=italic, color_rgb=color, east_asia="Microsoft YaHei")
            else:
                color, bold, italic = map_token_to_style(ttype, val, lang)
                r = p_cur.add_run(val)
                set_run_font(r, font_name="Consolas", size_pt=code_font_size, bold=bold, italic=italic, color_rgb=color, east_asia="Microsoft YaHei")

        for p in all_paras[:-1]:
            p.paragraph_format.keep_with_next = True

    # ---------------- BUILD DOCUMENT CONTENT ----------------
    # ==================== PAGE 1 ====================
    add_title(
        "Windows 桌面安装包构建与自动更新完整指南",
        "Next.js 纯前端 + Python FastAPI 伴侣进程 · 工业级发布与热更新全方案"
    )
    add_meta_bar()

    # Introduction
    add_body([
        ("把当前项目（", False, False, None),
        ("Next.js 前端", True, False, RGBColor(15, 44, 89)),
        (" + ", False, False, None),
        ("Python FastAPI 后端", True, False, RGBColor(15, 44, 89)),
        ("）改造成带有自动更新功能的 ", False, False, None),
        ("Windows 桌面安装包", True, False, RGBColor(15, 44, 89)),
        ("，分为以下 ", False, False, None),
        ("5 个完整阶段", True, False, RGBColor(180, 83, 9)),
        ("。下面为你梳理从准备环境、日常热更新开发、打包安装包、到最后发布自动更新的", False, False, None),
        ("端到端全流程", True, False, RGBColor(15, 44, 89)),
        ("：", False, False, None)
    ], space_after=2.5, space_before=2)

    # ---------------- 阶段一 ----------------
    add_h1("阶段一：项目环境与依赖准备", space_before=6, space_after=2)
    add_body([
        ("在当前项目根目录（", False, False, None),
        ("Tools/", True, False, RGBColor(15, 44, 89)),
        ("）下，安装桌面端外壳的核心依赖模块：", False, False, None)
    ], space_after=2)

    add_vscode_block(
        "bash",
        "BASH  —  终端命令：安装 Electron 与桌面端外壳依赖",
        """# 安装 electron 开发核心、打包构建器、以及支持自动更新的模块
npm install --save-dev electron electron-builder concurrently wait-on
npm install electron-updater tree-kill"""
    )

    add_body([
        ("安装的核心依赖库及其在架构中的关键职责说明如下：", False, False, None)
    ], space_before=2.5, space_after=1.5)

    add_bullet([
        ("electron-builder", True, False, RGBColor(15, 44, 89)),
        ("：负责把所有前端静态文件、主进程代码与 Python 独立运行时压缩封装成 Windows NSIS 一键安装程序（.exe）。", False, False, None)
    ])
    add_bullet([
        ("electron-updater", True, False, RGBColor(15, 44, 89)),
        ("：负责在客户端后台静默检查版本差分、下载更新包、SHA512 哈希校验并在用户重启时自动替换升级。", False, False, None)
    ])
    add_bullet([
        ("tree-kill", True, False, RGBColor(15, 44, 89)),
        ("：确保用户关闭软件或崩溃退出时，能够沿进程树递归彻底杀死后台常驻的 Python FastAPI 进程，杜绝后台僵尸进程与端口占用。", False, False, None)
    ])

    add_callout(
        "📌 图标资源准备规范 (resources/icon.ico)",
        [
            [
                ("准备一张 ", False, False, None),
                ("256×256 像素", True, False, RGBColor(30, 41, 59)),
                (" 的透明背景高清 Logo，使用转换工具生成标准的 ", False, False, None),
                ("icon.ico", True, False, RGBColor(15, 44, 89)),
                (" 多分辨率图标文件，保存在项目目录 ", False, False, None),
                ("resources/icon.ico", True, False, RGBColor(15, 44, 89)),
                ("。打包工具将自动注入至 Windows 快捷方式及应用程序图标中。", False, False, None)
            ]
        ],
        box_type="info"
    )

    # ==================== PAGE 2 ====================
    doc.add_page_break()

    # ---------------- 阶段二 ----------------
    add_h1("阶段二：Python 后端独立化（PyInstaller）", space_before=0, space_after=2)
    add_body([
        ("目标终端用户的电脑上绝大多数没有安装 Python 环境。因此，必须将后端的 FastAPI 框架、Uvicorn 服务以及 PyMuPDF 等算法依赖，编译打包成一个", False, False, None),
        ("完全独立的免安装目录", True, False, RGBColor(15, 44, 89)),
        ("，随客户端一并分发：", False, False, None)
    ], space_after=2)

    add_h2("1. 在 backend/ 虚拟环境中安装 PyInstaller", space_before=4, space_after=1.5)
    add_vscode_block(
        "bash",
        "BASH  —  终端命令：在后端环境中安装 PyInstaller",
        """pip install pyinstaller"""
    )

    add_h2("2. 执行 Python 后端打包命令", space_before=4, space_after=1.5)
    add_body([
        ("编写一个打包命令/脚本，将 ", False, False, None),
        ("app/main.py", True, False, RGBColor(15, 44, 89)),
        (" 打包为免安装目录：", False, False, None)
    ], space_after=2)

    add_vscode_block(
        "bash",
        "BASH  —  终端命令：PyInstaller 目录模式打包",
        """pyinstaller --noconfirm --onedir --name "omni-backend" --add-data "app;app" app/main.py"""
    )

    add_callout(
        "💡 打包产物与分发路径说明",
        [
            [
                ("打包完成后，会在 ", False, False, None),
                ("backend/dist/omni-backend/", True, False, RGBColor(15, 44, 89)),
                (" 目录下生成包含独立 Python 运行时及 ", False, False, None),
                ("omni-backend.exe", True, False, RGBColor(15, 44, 89)),
                (" 的免安装文件夹。", False, False, None)
            ],
            [
                ("这个文件夹后续会自动被 Electron 的 ", False, False, None),
                ("extraResources", True, False, RGBColor(15, 44, 89)),
                (" 机制一起打进最终安装包里，安装后内置于资源目录，无需终端用户手动安装 Python。", False, False, None)
            ]
        ],
        box_type="info"
    )

    add_callout(
        "⚙️ PyInstaller 关键参数解析",
        [
            [
                ("• --noconfirm", True, False, RGBColor(15, 44, 89)),
                ("：静默覆盖已有构建输出，防止因交互确认提示导致自动化打包脚本阻塞挂起。", False, False, None)
            ],
            [
                ("• --onedir", True, False, RGBColor(15, 44, 89)),
                ("：生成免安装目录格式。相比单文件 (--onefile) 启动耗时大幅缩减 90% 以上（无需每次运行向临时目录解压数百兆运行时）。", False, False, None)
            ],
            [
                ("• --add-data \"app;app\"", True, False, RGBColor(15, 44, 89)),
                ("：将 FastAPI 应用源码目录及内部静态依赖完整打包进独立执行目录。", False, False, None)
            ],
            [
                ("• --name \"omni-backend\"", True, False, RGBColor(15, 44, 89)),
                ("：指定输出可执行文件名，与后文 Electron 主进程中的子进程拉起逻辑精准匹配。", False, False, None)
            ]
        ],
        box_type="info"
    )

    # ==================== PAGE 3 ====================
    doc.add_page_break()

    add_h1("阶段三：编写 Electron 主进程（支持开发热更 & 生产静默运行）", space_before=0, space_after=2)
    add_body([
        ("在项目根目录下新建 ", False, False, None),
        ("electron/", True, False, RGBColor(15, 44, 89)),
        (" 目录，编写桌面端外壳生命周期与进程控制逻辑：", False, False, None)
    ], space_after=1.5)

    add_h2("1. electron/main.js（核心主进程控制器）", space_before=3, space_after=1.5)
    add_vscode_block(
        "javascript",
        "JS  —  electron/main.js",
        """const { app, BrowserWindow, ipcMain } = require('electron');
const path = require('path');
const { spawn } = require('child_process');
const treeKill = require('tree-kill');

const isDev = process.env.NODE_ENV === 'development';
let mainWindow = null;
let backendProcess = null;
const BACKEND_PORT = 18520; // 专属本地端口，防止和 8000 冲突

// 1. 启动 Python 伴侣进程
function startBackend() {
  const backendExe = isDev
    ? path.join(__dirname, '../backend/dist/omni-backend/omni-backend.exe')
    : path.join(process.resourcesPath, 'backend', 'omni-backend.exe');

  backendProcess = spawn(backendExe, ['--port', String(BACKEND_PORT)]);
}

// 2. 创建主窗口
function createWindow() {
  mainWindow = new BrowserWindow({
    width: 1280,
    height: 850,
    icon: path.join(__dirname, '../resources/icon.ico'),
    webPreferences: {
      preload: path.join(__dirname, 'preload.js'),
      nodeIntegration: false,
      contextIsolation: true,
    },
  });

  if (isDev) {
    // 【开发阶段】：直接连 Next.js 开发服务器，支持毫秒级热更新！
    mainWindow.loadURL('http://localhost:3000');
    mainWindow.webContents.openDevTools(); // 自动打开 F12 控制台
  } else {
    // 【生产发布】：加载打包好的本地 Next.js 静态文件
    mainWindow.loadFile(path.join(__dirname, '../frontend/out/index.html'));
  }
}

app.whenReady().then(() => {
  startBackend();
  createWindow();
});

// 3. 彻底退出处理：关闭窗口时杀死后台 Python
app.on('before-quit', () => {
  if (backendProcess && backendProcess.pid) {
    treeKill(backendProcess.pid, 'SIGKILL');
  }
});""",
        line_height_pt=9.6,
        code_font_size=8.3
    )

    # ==================== PAGE 4 ====================
    doc.add_page_break()

    add_h2("2. electron/preload.js（安全通讯通道）", space_before=0, space_after=1.5)
    add_vscode_block(
        "javascript",
        "JS  —  electron/preload.js",
        """const { contextBridge, ipcRenderer } = require('electron');

contextBridge.exposeInMainWorld('electronAPI', {
  checkForUpdates: () => ipcRenderer.invoke('check-for-updates'),
  startDownload: () => ipcRenderer.invoke('start-download'),
  quitAndInstall: () => ipcRenderer.invoke('quit-and-install'),
  onUpdateStatus: (callback) => ipcRenderer.on('update-status', (e, val) => callback(val)),
});"""
    )

    # ---------------- 阶段四 ----------------
    add_h1("阶段四：前端适配与自动更新交互", space_before=5, space_after=2)
    
    add_h2("1. Next.js 开启纯静态导出", space_before=3, space_after=1.5)
    add_body([
        ("修改 ", False, False, None),
        ("frontend/next.config.mjs", True, False, RGBColor(15, 44, 89)),
        ("，加上静态导出设置：", False, False, None)
    ], space_after=1.5)

    add_vscode_block(
        "javascript",
        "JS  —  frontend/next.config.mjs",
        """/** @type {import('next').NextConfig} */
const nextConfig = {
  output: 'export', // 导出静态 HTML/JS
  images: { unoptimized: true },
};

export default nextConfig;"""
    )

    add_h2("2. 前端请求地址兼容", space_before=3.5, space_after=1.5)
    add_body([
        ("前端请求 Python API 时，判断如果是桌面环境，自动请求 ", False, False, None),
        ("http://127.0.0.1:18520", True, False, RGBColor(15, 44, 89)),
        ("：", False, False, None)
    ], space_after=1.5)

    add_vscode_block(
        "typescript",
        "TS  —  API 地址自适应逻辑",
        """export const API_BASE_URL = typeof window !== 'undefined' && (window as any).electronAPI
  ? 'http://127.0.0.1:18520'
  : 'http://127.0.0.1:8000';"""
    )

    add_callout(
        "🔐 前端与桌面端通信架构设计要点",
        [
            [
                ("• 上下文隔离保障", True, False, RGBColor(15, 44, 89)),
                ("：关闭 nodeIntegration 并启用 contextIsolation，仅通过 preload.js 向渲染进程暴露受控 API，杜绝 XSS 漏洞直接执行系统命令的安全隐患。", False, False, None)
            ],
            [
                ("• 纯静态导出架构", True, False, RGBColor(15, 44, 89)),
                ("：Next.js 输出纯粹的 HTML/JS/CSS，无需在客户端运行 Node.js 服务端渲染，极大降低内存占用并实现秒开体验。", False, False, None)
            ]
        ],
        box_type="info"
    )

    # ==================== PAGE 5 ====================
    doc.add_page_break()

    add_h2("3. 主进程接入 electron-updater（监听更新事件）", space_before=0, space_after=1.5)
    add_body([
        ("在 ", False, False, None),
        ("electron/updater.js", True, False, RGBColor(15, 44, 89)),
        (" 中配置自动更新生命周期事件与 IPC 处理句柄：", False, False, None)
    ], space_after=1.5, keep_with_next=True)

    add_vscode_block(
        "javascript",
        "JS  —  electron/updater.js",
        """const { autoUpdater } = require('electron-updater');
const { ipcMain } = require('electron');

function setupAutoUpdater(win) {
  autoUpdater.autoDownload = false; // 用户点击后再下载

  // 监听到新版本
  autoUpdater.on('update-available', (info) => {
    win.webContents.send('update-status', { step: 'available', version: info.version });
  });

  // 下载进度 (0% ~ 100%)
  autoUpdater.on('download-progress', (progress) => {
    win.webContents.send('update-status', { step: 'downloading', percent: progress.percent });
  });

  // 下载完毕，准备安装
  autoUpdater.on('update-downloaded', () => {
    win.webContents.send('update-status', { step: 'ready' });
  });

  ipcMain.handle('check-for-updates', () => autoUpdater.checkForUpdates());
  ipcMain.handle('start-download', () => autoUpdater.downloadUpdate());
  ipcMain.handle('quit-and-install', () => autoUpdater.quitAndInstall());
}

module.exports = { setupAutoUpdater };""",
        line_height_pt=9.6,
        code_font_size=8.3
    )

    add_callout(
        "🔄 自动更新全生命周期原理解析",
        [
            [
                ("• 1. 检查更新 (check-for-updates)", True, False, RGBColor(15, 44, 89)),
                ("：主进程请求 GitHub Releases 接口获取 latest.yml 配置文件，比对本地与远端语义化版本号。", False, False, None)
            ],
            [
                ("• 2. 静默下载 (start-download)", True, False, RGBColor(15, 44, 89)),
                ("：通过 download-progress 实时向渲染进程派发下载进度（0%~100%），前端可在 UI 呈现平滑进度条。", False, False, None)
            ],
            [
                ("• 3. 完整性校验与就绪 (update-downloaded)", True, False, RGBColor(15, 44, 89)),
                ("：更新包下载完毕后自动比对 SHA512 哈希指纹，防止篡改与网络损坏，就绪后提示用户重启。", False, False, None)
            ],
            [
                ("• 4. 退出并安装 (quit-and-install)", True, False, RGBColor(15, 44, 89)),
                ("：主进程退出当前应用，静默调用 NSIS 安装程序完成原位覆盖替换，并自动拉起全新版本。", False, False, None)
            ]
        ],
        box_type="info"
    )

    # ==================== PAGE 6 ====================
    doc.add_page_break()

    # ---------------- 阶段五 ----------------
    add_h1("阶段五：配置打包与一键发布", space_before=0, space_after=2)
    add_h2("1. 在根目录 package.json 配置打包规则", space_before=3, space_after=1.5)
    
    add_vscode_block(
        "json",
        "JSON  —  package.json (根目录完整打包与更新分发配置)",
        """{
  "name": "xc-omnibox",
  "version": "1.0.0",
  "main": "electron/main.js",
  "scripts": {
    "dev:electron": "concurrently \\"npm --prefix frontend run dev\\" \\"wait-on http://localhost:3000 && electron .\\"",
    "build:frontend": "npm --prefix frontend run build",
    "build:backend": "cd backend && pyinstaller --noconfirm --onedir --name omni-backend app/main.py",
    "dist": "npm run build:frontend && npm run build:backend && electron-builder"
  },
  "build": {
    "appId": "com.xc.omnibox",
    "productName": "XC 万象箱",
    "directories": {
      "output": "release"
    },
    "extraResources": [
      {
        "from": "backend/dist/omni-backend",
        "to": "backend"
      }
    ],
    "publish": [
      {
        "provider": "github",
        "owner": "LEESC88",
        "repo": "XC_OmniBox"
      }
    ],
    "win": {
      "target": ["nsis"],
      "icon": "resources/icon.ico"
    },
    "nsis": {
      "oneClick": true,
      "perMachine": false,
      "createDesktopShortcut": true,
      "createStartMenuShortcut": true,
      "shortcutName": "XC 万象箱"
    }
  }
}""",
        line_height_pt=9.6,
        code_font_size=8.3
    )

    add_callout(
        "📦 package.json 关键配置字段解析",
        [
            [
                ("• extraResources", True, False, RGBColor(15, 44, 89)),
                ("：将 PyInstaller 打包好的 Python 后端文件夹完整复制到安装目录的 resources/ 资源文件夹下，客户端直接调用，免除外部依赖。", False, False, None)
            ],
            [
                ("• publish (provider: github)", True, False, RGBColor(15, 44, 89)),
                ("：声明版本发布的远程仓库，electron-updater 会自动请求该 GitHub 仓库的 Releases 接口对比 latest.yml。", False, False, None)
            ],
            [
                ("• nsis (oneClick: true)", True, False, RGBColor(15, 44, 89)),
                ("：启用一键极速静默安装，无需向导式弹窗，自动创建桌面及开始菜单快捷方式，带来现代主流软件的用户体验。", False, False, None)
            ]
        ],
        box_type="info"
    )

    # ==================== PAGE 7 ====================
    doc.add_page_break()

    # ---------------- 以后你的两套工作流 ----------------
    add_h1("以后你的两套工作流", space_before=0, space_after=2)
    
    add_h2("工作流 A：日常写代码（99% 的时间）", space_before=3, space_after=1.5)
    add_vscode_block(
        "bash",
        "BASH  —  终端命令：启动热更新开发环境",
        """npm run dev:electron"""
    )

    add_bullet([
        ("一行命令同时拉起 Next.js 和桌面测试窗口。", False, False, None)
    ])
    add_bullet([
        ("在 VS Code 里改代码，桌面窗口直接热刷新，按 ", False, False, None),
        ("Ctrl + S", True, False, RGBColor(180, 83, 9)),
        (" 立即看到效果。", False, False, None)
    ])

    add_h2("工作流 B：发布新版本（要给别人更新时）", space_before=4, space_after=1.5)
    
    add_bullet([
        ("1. 把 ", False, False, None),
        ("package.json", True, False, RGBColor(15, 44, 89)),
        (" 里的版本号改成 ", False, False, None),
        ("1.0.1", True, False, RGBColor(180, 83, 9)),
        ("。", False, False, None)
    ])
    add_bullet([
        ("2. 运行打包命令：", False, False, None)
    ])

    add_vscode_block(
        "bash",
        "BASH  —  终端命令：打包发布生产安装包",
        """npm run dist"""
    )

    add_bullet([
        ("3. 1~2 分钟后，", False, False, None),
        ("release/", True, False, RGBColor(15, 44, 89)),
        (" 文件夹下会生成：", False, False, None)
    ])
    add_bullet([
        ("XC 万象箱 Setup 1.0.1.exe", True, False, RGBColor(15, 44, 89)),
        ("（安装包）", False, False, None)
    ], level=1)
    add_bullet([
        ("latest.yml", True, False, RGBColor(15, 44, 89)),
        ("（版本更新配置文件）", False, False, None)
    ], level=1)

    add_bullet([
        ("4. 把这两个文件上传到你 GitHub 仓库的 ", False, False, None),
        ("Releases", True, False, RGBColor(15, 44, 89)),
        (" 页面发布。", False, False, None)
    ])

    add_callout(
        "🚀 自动更新升级效果",
        [
            [
                ("所有用户打开电脑上的老版本软件时，前端就会收到通知，点击“下载更新”，下载完毕点击“重启”，软件自动完成替换并升级到 ", False, False, None),
                ("1.0.1", True, False, RGBColor(22, 101, 52)),
                ("！", False, False, None)
            ]
        ],
        box_type="success"
    )

    add_callout(
        "🎯 方案结语",
        [
            [
                ("这就是完整的工业级桌面端方案。整个骨架非常清晰，我们可以一步一步来实施。", False, False, RGBColor(30, 41, 59))
            ]
        ],
        box_type="info"
    )

    output_path = "Application_exe building.docx"
    doc.save(output_path)
    print(f"Successfully saved {output_path}")

if __name__ == "__main__":
    create_document()
